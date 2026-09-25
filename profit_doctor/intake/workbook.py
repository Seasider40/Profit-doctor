from __future__ import annotations
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any
from openpyxl import load_workbook

DOMAIN_RULES = [
 ('D04_AR', ('accounts receivable','ar customer','debtor','customer ledger')),
 ('D05_AP', ('accounts payable','ap supplier','creditor','supplier ledger')),
 ('D11_WORKFORCE', ('staff cost','employee','payroll')),
 ('D16_OPERATIONS', ('capacity','utilisation','work centre')),
 ('D03_TB_GL', ('trial balance','account code')),
 ('D01_PNL', ('management accounts','gross profit','ebitda','revenue')),
 ('CONTROL_EVIDENCE', ('reconciliation','control account')),
 ('CONTEXT_CONFIG', ('assumption','overview')),
]

def _norm(v): return str(v or '').strip().lower()

def _header_row(ws):
    best=(1,0)
    for r in range(1,min(ws.max_row or 0,12)+1):
        vals=[ws.cell(r,c).value for c in range(1,(ws.max_column or 0)+1)]
        score=sum(v is not None and str(v).strip()!='' for v in vals)
        if score>best[1]: best=(r,score)
    return best[0]

def classify_sheet(name:str, headers:list[Any], sample:list[list[Any]]|None=None):
    text=' '.join([_norm(name)]+[_norm(x) for x in headers])
    scores=[]
    for domain,terms in DOMAIN_RULES:
        hits=sum(1 for t in terms if t in text)
        if hits: scores.append((hits,domain))
    scores.sort(reverse=True)
    return scores[0][1] if scores else 'UNKNOWN'

def profile_workbook(path:str|Path):
    p=Path(path)
    wbf=load_workbook(p,data_only=False,read_only=False)
    wbv=load_workbook(p,data_only=True,read_only=False)
    sheets=[]
    for ws in wbf.worksheets:
        hr=_header_row(ws)
        headers=[ws.cell(hr,c).value for c in range(1,(ws.max_column or 0)+1)]
        formulas=sum(1 for row in ws.iter_rows() for cell in row if isinstance(cell.value,str) and cell.value.startswith('='))
        values_ws=wbv[ws.title]
        cached_missing=sum(1 for row in ws.iter_rows() for cell in row if isinstance(cell.value,str) and cell.value.startswith('=') and values_ws[cell.coordinate].value is None)
        domain=classify_sheet(ws.title,headers)
        sheets.append({'sheet':ws.title,'rows':ws.max_row or 0,'columns':ws.max_column or 0,'header_row':hr,'headers':[str(x) if x is not None else '' for x in headers], 'domain':domain,'formula_cells':formulas,'formula_values_missing':cached_missing})
    return {'file':p.name,'sheet_count':len(sheets),'sheets':sheets}

def _find_sheet(wb, needles):
    for ws in wb.worksheets:
        n=_norm(ws.title)
        if any(x in n for x in needles): return ws
    return None

def _find_tb_amount(ws, account_terms):
    # supports debit/credit TBs; returns signed debit-positive balance for matching row
    hr=_header_row(ws); headers=[_norm(ws.cell(hr,c).value) for c in range(1,ws.max_column+1)]
    name_col=next((i+1 for i,h in enumerate(headers) if h in ('account','account name')),2)
    debit_col=next((i+1 for i,h in enumerate(headers) if 'debit' in h),None)
    credit_col=next((i+1 for i,h in enumerate(headers) if 'credit' in h),None)
    for r in range(hr+1,ws.max_row+1):
        name=_norm(ws.cell(r,name_col).value)
        if any(t in name for t in account_terms):
            d=float(ws.cell(r,debit_col).value or 0) if debit_col else 0
            c=float(ws.cell(r,credit_col).value or 0) if credit_col else 0
            return d-c
    return None

def _sum_col(ws, header_terms):
    hr=_header_row(ws); headers=[_norm(ws.cell(hr,c).value) for c in range(1,ws.max_column+1)]
    col=next((i+1 for i,h in enumerate(headers) if any(t in h for t in header_terms)),None)
    if not col:return None
    total=0.0; seen=False
    # Sum detail rows only.  SME schedules commonly include a final Total row;
    # including it would double-count the entire subledger.
    label_col=1
    for r in range(hr+1,ws.max_row+1):
        label=_norm(ws.cell(r,label_col).value)
        if 'total' in label or label == 'subtotal':
            continue
        v=ws.cell(r,col).value
        if isinstance(v,(int,float)): total+=float(v); seen=True
    return total if seen else None

def reconcile_workbook(path:str|Path, tolerance=1000.0):
    wb=load_workbook(path,data_only=True,read_only=False)
    tb=_find_sheet(wb,('trial balance',)); ar=_find_sheet(wb,('ar customer','customers')); ap=_find_sheet(wb,('ap supplier',))
    out=[]
    if tb and ar:
        gl=_find_tb_amount(tb,('trade receiv','accounts receiv'))
        sub=_sum_col(ar,('ar balance','receivable balance'))
        if gl is not None and sub is not None:
            diff=gl-sub; out.append({'control':'AR_TO_TB','gl':gl,'supporting':sub,'difference':diff,'status':'PASS' if abs(diff)<=tolerance else 'FAIL'})
    if tb and ap:
        gl=abs(_find_tb_amount(tb,('trade payable','accounts payable')) or 0)
        sub=_sum_col(ap,('ap balance','payable balance'))
        if gl and sub is not None:
            diff=gl-sub; out.append({'control':'AP_TO_TB','gl':gl,'supporting':sub,'difference':diff,'status':'PASS' if abs(diff)<=tolerance else 'FAIL'})
    return out

# v2.24 semantic mapping / reconciliation / canonical hand-off foundation.
SEMANTIC_TERMS={
 'entity.customer':('customer','client'), 'entity.supplier':('supplier','vendor'), 'entity.employee':('employee','staff'),
 'entity.work_centre':('work centre','work center'), 'metric.revenue':('annual sales','sales (£)','revenue'),
 'metric.cogs':('cogs','cost of sales'), 'metric.gross_profit':('gross profit',), 'metric.gross_margin_pct':('gross margin %','margin %'),
 'metric.ar_balance':('ar balance','receivable balance'), 'metric.ap_balance':('ap balance','payable balance'),
 'metric.debtor_days':('debtor days','dso'), 'metric.purchases':('annual purchases','purchases'),
 'metric.fte':('fte','fte / machines'), 'metric.total_staff_cost':('total cost',),
 'metric.available_hours':('annual available hrs','available hrs'), 'metric.practical_capacity_hours':('practical capacity hrs',),
 'metric.planned_load_hours':('planned load hrs','budget load hrs'), 'metric.actual_productive_hours':('actual productive hrs',),
 'metric.utilisation_pct':('actual utilisation %','planned utilisation %','budget utilisation %'),
 'account.code':('account code','code'), 'account.name':('account name','account'), 'account.debit':('debit','debit (£)'), 'account.credit':('credit','credit (£)'),
}

def _semantic(header):
    h=_norm(header)
    best=None
    for semantic,terms in SEMANTIC_TERMS.items():
        for t in terms:
            if h==t or t in h:
                score=1.0 if h==t else .85
                if best is None or score>best[0]: best=(score,semantic)
    return best

def _unit(header, values):
    h=_norm(header)
    if '%' in h or 'percent' in h: return 'PERCENT'
    if '£' in h or 'gbp' in h or any(x in h for x in ('sales','cogs','cost','balance','purchases','profit')): return 'GBP'
    if 'hrs' in h or 'hours' in h: return 'HOURS'
    if 'days' in h: return 'DAYS'
    if 'fte' in h: return 'FTE'
    return 'TEXT' if not any(isinstance(v,(int,float)) for v in values if v is not None) else 'NUMBER'

def semantic_map_workbook(path:str|Path):
    wb=load_workbook(path,data_only=True,read_only=False); out=[]
    for ws in wb.worksheets:
        hr=_header_row(ws); headers=[ws.cell(hr,c).value for c in range(1,ws.max_column+1)]
        domain=classify_sheet(ws.title,headers); cols=[]
        for c,h in enumerate(headers,1):
            vals=[ws.cell(r,c).value for r in range(hr+1,min(ws.max_row,hr+20)+1)]
            sem=_semantic(h)
            cols.append({'column':c,'header':str(h or ''),'semantic':sem[1] if sem else None,'confidence':sem[0] if sem else 0.0,'unit':_unit(h,vals),'requires_confirmation':not sem or sem[0]<.9})
        mapped=sum(1 for x in cols if x['semantic'])
        out.append({'sheet':ws.title,'domain':domain,'columns':cols,'mapping_confidence':round(mapped/max(1,len([h for h in headers if h not in (None,'')])),3),'requires_human_confirmation':domain=='UNKNOWN' or any(x['requires_confirmation'] for x in cols if x['header'])})
    return out

def _tb_value(tb, terms):
    v=_find_tb_amount(tb,terms); return None if v is None else float(v)

def generic_reconciliations(path:str|Path,tolerance=1000.0):
    wb=load_workbook(path,data_only=True,read_only=False); out=reconcile_workbook(path,tolerance)
    tb=_find_sheet(wb,('trial balance',))
    # Inventory: compare GL to any reconciliation/control evidence only as supporting evidence, never trust its status/formula.
    rec=_find_sheet(wb,('reconciliation',))
    if tb and rec:
        # Aggregate all inventory/stock TB accounts; do not stop at the first component account.
        thr=_header_row(tb); th=[_norm(tb.cell(thr,c).value) for c in range(1,tb.max_column+1)]
        tn=next((i+1 for i,h in enumerate(th) if h in ('account','account name')),2); td=next((i+1 for i,h in enumerate(th) if 'debit' in h),None); tc=next((i+1 for i,h in enumerate(th) if 'credit' in h),None)
        inv_parts=[]
        for rr in range(thr+1,tb.max_row+1):
            nm=_norm(tb.cell(rr,tn).value)
            if 'inventory' in nm or 'stock' in nm:
                d=float(tb.cell(rr,td).value or 0) if td else 0; c=float(tb.cell(rr,tc).value or 0) if tc else 0; inv_parts.append(d-c)
        inv=sum(inv_parts) if inv_parts else None
        hr=_header_row(rec)
        for r in range(hr+1,rec.max_row+1):
            label=_norm(rec.cell(r,1).value)
            if 'inventory' in label and inv is not None:
                supporting=rec.cell(r,3).value
                if isinstance(supporting,(int,float)):
                    diff=inv-float(supporting); out.append({'control':'INVENTORY_TO_SUPPORT','gl':inv,'supporting':float(supporting),'difference':diff,'status':'PASS' if abs(diff)<=tolerance else 'FAIL','basis':'independently recalculated from TB and supporting amount; embedded status ignored'})
            if ('debt' in label or 'borrow' in label) and tb:
                supporting=rec.cell(r,3).value
                if isinstance(supporting,(int,float)):
                    # debt can span multiple TB lines; aggregate common borrowing terms.
                    total=0.0; found=False
                    thr=_header_row(tb); headers=[_norm(tb.cell(thr,c).value) for c in range(1,tb.max_column+1)]
                    name_col=next((i+1 for i,h in enumerate(headers) if h in ('account','account name')),2); debit_col=next((i+1 for i,h in enumerate(headers) if 'debit' in h),None); credit_col=next((i+1 for i,h in enumerate(headers) if 'credit' in h),None)
                    for rr in range(thr+1,tb.max_row+1):
                        nm=_norm(tb.cell(rr,name_col).value)
                        if any(t in nm for t in ('loan','debt','overdraft','borrowing')):
                            d=float(tb.cell(rr,debit_col).value or 0) if debit_col else 0; c=float(tb.cell(rr,credit_col).value or 0) if credit_col else 0
                            total+=abs(d-c); found=True
                    if found:
                        diff=total-float(supporting); out.append({'control':'DEBT_TO_SUPPORT','gl':total,'supporting':float(supporting),'difference':diff,'status':'PASS' if abs(diff)<=tolerance else 'FAIL','basis':'independently recalculated from TB borrowing accounts; embedded status ignored'})
    return out

def canonical_extract(path:str|Path):
    """Return normalized evidence records suitable for downstream domain adapters.
    Does not silently invent invoice/transaction detail that the workbook does not contain.
    """
    wb=load_workbook(path,data_only=True,read_only=False); result={}
    for ws in wb.worksheets:
        hr=_header_row(ws); headers=[ws.cell(hr,c).value for c in range(1,ws.max_column+1)]; domain=classify_sheet(ws.title,headers)
        sm=semantic_map_workbook(path); smrow=next(x for x in sm if x['sheet']==ws.title); bycol={x['column']:x for x in smrow['columns'] if x['semantic']}
        if domain in ('D04_AR','D05_AP','D11_WORKFORCE','D16_OPERATIONS'):
            records=[]
            for r in range(hr+1,ws.max_row+1):
                label=_norm(ws.cell(r,1).value)
                if not label or 'total' in label or label=='subtotal': continue
                rec={'source_sheet':ws.title,'source_row':r}
                for c,m in bycol.items():
                    rec[m['semantic']]=ws.cell(r,c).value
                if len(rec)>2: records.append(rec)
            result.setdefault(domain,[]).extend(records)
    return result

def intake_assessment(path:str|Path):
    profile=profile_workbook(path); mappings=semantic_map_workbook(path); recons=generic_reconciliations(path); canonical=canonical_extract(path)
    available=sorted({s['domain'] for s in profile['sheets'] if s['domain'].startswith('D')})
    failed=[r for r in recons if r['status']=='FAIL']
    return {'profile':profile,'mappings':mappings,'reconciliations':recons,'canonical_domains':{k:len(v) for k,v in canonical.items()},'available_domains':available,'integrity_state':'MATERIALLY_CONSTRAINED' if failed else 'USABLE','failed_controls':[r['control'] for r in failed]}
