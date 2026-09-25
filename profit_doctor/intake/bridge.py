from __future__ import annotations
from pathlib import Path
from tempfile import TemporaryDirectory
from datetime import datetime, timezone
import csv, re, uuid
from decimal import Decimal
from openpyxl import load_workbook
from profit_doctor.core.db import connect
from profit_doctor.ingestion.accounting import ingest_accounting_file
from profit_doctor.intake.workbook import _header_row, _norm, semantic_map_workbook, generic_reconciliations, intake_assessment, canonical_extract
from profit_doctor.calc.primitive_engine import run_primitive_engine
from profit_doctor.diagnostic.engine import run_diagnostic_engine
from profit_doctor.reasoning.engine import run_reasoning_engine
from profit_doctor.economic.engine import run_economic_engine
from profit_doctor.management.attention import build_management_attention
from profit_doctor.management.opportunity_register import build_action_opportunity_register

now=lambda: datetime.now(timezone.utc).isoformat()
def _id(p): return f'{p}_{uuid.uuid4().hex}'

def _year_from_title(s, default=2026):
    m=re.search(r'(20\d{2})',s or '')
    return int(m.group(1)) if m else default

def _ma_rows(path):
    wb=load_workbook(path,data_only=True); ws=next((w for w in wb.worksheets if 'management accounts' in _norm(w.title)),None)
    if not ws: return []
    hr=_header_row(ws); year=_year_from_title(str(ws.cell(1,1).value or ''))
    months=['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec']
    h=[str(ws.cell(hr,c).value or '') for c in range(1,ws.max_column+1)]
    monthcols=[(c,m) for c,m in enumerate(h,1) if m in months]
    aliases={'revenue':'REV','sales':'REV','turnover':'REV','cost of sales':'COGS','cogs':'COGS','gross profit':'GP','ebitda':'EBITDA'}
    out=[]
    for r in range(hr+1,ws.max_row+1):
        label=_norm(ws.cell(r,1).value); code=None
        for a,k in aliases.items():
            if label==a: code=k; break
        if not code: continue
        for c,m in monthcols:
            v=ws.cell(r,c).value
            if isinstance(v,(int,float)):
                mi=months.index(m)+1
                import calendar
                d=calendar.monthrange(year,mi)[1]
                # workbook declares £000 unless stated
                out.append({'period_end':f'{year}-{mi:02d}-{d:02d}','line_code':code,'line_name':label.title(),'amount':float(v)*1000})
    return out

def _tb_rows(path):
    wb=load_workbook(path,data_only=True); ws=next((w for w in wb.worksheets if 'trial balance' in _norm(w.title)),None)
    if not ws:return []
    hr=_header_row(ws); year=_year_from_title(str(ws.cell(1,1).value or '')); headers=[_norm(ws.cell(hr,c).value) for c in range(1,ws.max_column+1)]
    def col(terms,default=None): return next((i+1 for i,h in enumerate(headers) if any(t==h or t in h for t in terms)),default)
    cc=col(('account code','code'),1); nc=next((i+1 for i,h in enumerate(headers) if h=='account name' or h=='account'),2); cat=col(('category',),3); dc=col(('debit',),4); cr=col(('credit',),5)
    out=[]
    for r in range(hr+1,ws.max_row+1):
        code=ws.cell(r,cc).value; name=ws.cell(r,nc).value
        if code in (None,'') or name in (None,''): continue
        d=ws.cell(r,dc).value or 0; c=ws.cell(r,cr).value or 0
        if not isinstance(d,(int,float)) or not isinstance(c,(int,float)): continue
        out.append({'period_end':f'{year}-12-31','account_code':str(code),'account_name':str(name),'account_type':str(ws.cell(r,cat).value or 'UNKNOWN'),'debit':float(d),'credit':float(c)})
    return out

def _bs_from_tb(tb):
    # Conservative control-balance extraction. Broad tokens such as 'bank', 'debtor'
    # or 'payable' can silently pull debt, tax or other-debtor accounts into cash/AR/AP.
    # Unknown-workbook intake therefore prefers explicit trade/control account semantics.
    def classify(name):
        n=_norm(name)
        if any(t in n for t in ('trade receivable','accounts receivable','trade debtor')): return 'AR'
        if any(t in n for t in ('trade payable','accounts payable','trade creditor')): return 'AP'
        if any(t in n for t in ('cash at bank','cash and cash equivalent','cash in bank')) or n=='cash': return 'CASH'
        if any(t in n for t in ('inventory','stock')): return 'INVENTORY'
        return None
    vals={k:Decimal('0') for k in ('AR','AP','CASH','INVENTORY')}; seen={k:False for k in vals}
    for r in tb:
        k=classify(r['account_name'])
        if not k: continue
        bal=Decimal(str(r['debit']))-Decimal(str(r['credit']))
        vals[k]+=bal; seen[k]=True
    pe=tb[0]['period_end'] if tb else None; out=[]
    for k,v in vals.items():
        if seen[k]: out.append({'period_end':pe,'line_code':k,'line_name':k,'amount':float(abs(v) if k=='AP' else v)})
    return out

def _write_csv(path, rows, fields):
    with open(path,'w',newline='',encoding='utf-8') as f:
        w=csv.DictWriter(f,fieldnames=fields); w.writeheader(); w.writerows(rows)

def _insert_workforce(con,client,path):
    data=canonical_extract(path).get('D11_WORKFORCE',[]); n=0
    for i,r in enumerate(data,1):
        emp=r.get('entity.employee') or f'EMP_{i:03d}'; fte=r.get('metric.fte') or 1; total=Decimal(str(r.get('metric.total_staff_cost') or 0))
        con.execute('INSERT OR REPLACE INTO workforce_snapshot VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',(_id('wf'),client,'2026-12-31',str(emp),r.get('Department') or r.get('department'),r.get('Role') or r.get('role'),str(fte),str(total),'0','0','0','0',None,None,'Unknown-workbook semantic mapping; total staff cost stored as base_salary surrogate because component mapping is incomplete'))
        n+=1
    con.commit(); return n

def execute_unknown_workbook(path, db_path=':memory:', client_id='UWB_CLIENT'):
    con=connect(db_path); run=_id('run')
    con.execute('INSERT INTO client VALUES (?,?,?,?,?)',(client_id,Path(path).stem,'GBP','PRODUCT_DISTRIBUTION',now()))
    con.execute('INSERT INTO engine_run VALUES (?,?,?,?,?,?,?,?,?)',(run,client_id,'UNKNOWN_WORKBOOK_QUALIFICATION',now(),None,'RUNNING',None,None,'2.28'))
    pnl=_ma_rows(path); tb=_tb_rows(path); bs=_bs_from_tb(tb)
    with TemporaryDirectory() as td:
        td=Path(td); store=td/'store'; store.mkdir()
        if pnl:
            p=td/'pnl.csv'; _write_csv(p,pnl,['period_end','line_code','line_name','amount']); ingest_accounting_file(con,client_id,run,p,'D01_PNL',store)
        if bs:
            p=td/'bs.csv'; _write_csv(p,bs,['period_end','line_code','line_name','amount']); ingest_accounting_file(con,client_id,run,p,'D02_BALANCE_SHEET',store)
        if tb:
            p=td/'tb.csv'; _write_csv(p,tb,['period_end','account_code','account_name','account_type','debit','credit']); ingest_accounting_file(con,client_id,run,p,'D03_TRIAL_BALANCE',store)
        workforce=_insert_workforce(con,client_id,path)
        # Persist intake controls as integrity evidence. Do not manufacture invoice-level AR/AP from aggregate schedules.
        controls=generic_reconciliations(path)
        for x in controls:
            con.execute('INSERT INTO reconciliation VALUES (?,?,?,?,?,?,?,?,?,?,?,?)',(_id('rec'),run,client_id,'UWB_'+x['control'],'ACCOUNTING_CONTROL','SUPPORTING_SCHEDULE',str(x.get('gl')),str(x.get('supporting')),str(x.get('difference')),('RECONCILED' if x['status']=='PASS' else 'FAILED'),x.get('basis'),now()))
        con.commit()
        run_primitive_engine(con,run,client_id)
        diagnostics=run_diagnostic_engine(con,run,client_id)
        aggregate_handoff=_aggregate_commercial_handoff(con,run,client_id,path)
        diagnostics.update(aggregate_handoff['tests'])
        reasoning=run_reasoning_engine(con,run,client_id)
        attention=build_management_attention(con,run,client_id)
        economic=run_economic_engine(con,run,client_id)
        opportunity_register=build_action_opportunity_register(con,run,client_id)
    executions=con.execute('SELECT test_id,execution_status,eligibility_state,signal_count,limitation FROM test_execution WHERE run_id=? ORDER BY test_id',(run,)).fetchall()
    completed=sum(1 for r in executions if r['execution_status']=='COMPLETED'); notrun=sum(1 for r in executions if r['execution_status']!='COMPLETED')
    signals=con.execute('SELECT count(*) FROM signal WHERE run_id=?',(run,)).fetchone()[0]
    findings=con.execute('SELECT count(*) FROM finding_version WHERE run_id=?',(run,)).fetchone()[0]
    con.execute("UPDATE engine_run SET status='COMPLETED',completed_at=? WHERE run_id=?",(now(),run)); con.commit()
    result={'run_id':run,'intake':intake_assessment(path),'canonical':{'pnl_rows':len(pnl),'bs_rows':len(bs),'tb_rows':len(tb),'workforce_rows':workforce},'controls':controls,'aggregate_handoff':aggregate_handoff,'diagnostics':{'completed':completed,'not_run':notrun,'total':len(executions),'signals':signals,'findings':findings},'reasoning':reasoning,'management_attention':attention,'economic':economic,'opportunity_register':opportunity_register,'executions':[dict(r) for r in executions]}
    con.close(); return result

# v2.26 evidence-granularity-aware commercial canonicalisation.
def _replace_execution(con, run, test_id):
    ids=[r['test_execution_id'] for r in con.execute('SELECT test_execution_id FROM test_execution WHERE run_id=? AND test_id=?',(run,test_id)).fetchall()]
    for x in ids:
        sids=[r['signal_id'] for r in con.execute('SELECT signal_id FROM signal WHERE test_execution_id=?',(x,)).fetchall()]
        for sid in sids: con.execute('DELETE FROM diagnostic_lineage WHERE signal_id=?',(sid,))
        con.execute('DELETE FROM signal WHERE test_execution_id=?',(x,)); con.execute('DELETE FROM test_execution WHERE test_execution_id=?',(x,))
    con.execute('DELETE FROM test_eligibility WHERE run_id=? AND test_id=?',(run,test_id,)); con.commit()

def _elig(con,run,client,test,method,state,lim=None):
    con.execute('INSERT OR REPLACE INTO test_eligibility VALUES (?,?,?,?,?,?,?,?,?,?,?,?)',(_id('elig'),run,client,test,'APPLICABLE',state,method,None,None,None,lim,now())); con.commit()

def _aggregate_commercial_handoff(con,run,client,path):
    from profit_doctor.diagnostic.engine import _execute, D
    data=canonical_extract(path); results={}
    customers=data.get('D04_AR',[]); suppliers=data.get('D05_AP',[]); ops=data.get('D16_OPERATIONS',[])
    # Annual customer summaries legitimately support concentration, Contribution 0 profitability and Pareto distribution.
    usable=[]
    for r in customers:
        name=r.get('entity.customer'); rev=r.get('metric.revenue'); gp=r.get('metric.gross_profit'); gm=r.get('metric.gross_margin_pct')
        if name and isinstance(rev,(int,float)):
            rev=D(rev)
            if isinstance(gp,(int,float)): gp=D(gp)
            elif isinstance(gm,(int,float)): gp=rev*D(gm) if abs(D(gm))<=D('2') else rev*D(gm)/D('100')
            usable.append((str(name),rev,gp,r['source_sheet'],r['source_row']))
    if usable:
        total=sum((x[1] for x in usable),D('0')); gps=sum((x[2] for x in usable if x[2] is not None),D('0'))
        shares=sorted(usable,key=lambda x:x[1],reverse=True); top5=sum((x[1] for x in shares[:5]),D('0')); hhi=sum(((x[1]/total)**2 for x in usable),D('0'))*D('10000') if total else D('0')
        _replace_execution(con,run,'CUS-01'); _elig(con,run,client,'CUS-01','AGGREGATE_CUSTOMER_ANNUAL','PARTIAL-A','Annual customer summary supports concentration; no transaction cadence, retention or invoice behaviour inferred')
        sig=[{'type':'TOP_CUSTOMER_CONCENTRATION','entity_type':'CUSTOMER','entity_id':shares[0][0],'observed':shares[0][1]/total*100,'unit':'PERCENT','materiality':'HIGH' if shares[0][1]/total>=D('.25') else 'MEDIUM','evidence':f'Annual customer summary: largest customer contributes {shares[0][1]} of {total} captured revenue. Dependency exposure only; no loss probability inferred.','primitive':None,'lineage':[('WORKBOOK_ROW',f'{shares[0][3]}:{shares[0][4]}','aggregate customer summary')]},{'type':'TOP5_CONCENTRATION','observed':top5/total*100,'unit':'PERCENT','materiality':'MEDIUM','evidence':f'Top five customers contribute {top5/total*100}% of captured annual customer revenue.','primitive':None,'lineage':[('WORKBOOK_SHEET',shares[0][3],'aggregate customer summary')]},{'type':'CUSTOMER_HHI','observed':hhi,'unit':'INDEX','materiality':'INFORMATIONAL','evidence':f'Customer revenue HHI is {hhi} from annual aggregate customer evidence.','primitive':None,'lineage':[('WORKBOOK_SHEET',shares[0][3],'aggregate customer summary')]}]
        results['CUS-01']=_execute(con,run,client,'CUS-01','AGGREGATE_CUSTOMER_ANNUAL','PARTIAL-A','COMPLETED',sig,'Aggregate annual evidence; no time-series customer movement inferred')
        if any(x[2] is not None for x in usable):
            _replace_execution(con,run,'CUS-02'); _elig(con,run,client,'CUS-02','AGGREGATE_CUSTOMER_MARGIN','PARTIAL-A','Annual customer revenue/margin supports Contribution 0 economics; no cost-to-serve inferred')
            sig=[]
            for name,rev,gp,sh,row in usable:
                if gp is None: continue
                sig.append({'type':'CUSTOMER_PROFITABILITY','entity_type':'CUSTOMER','entity_id':name,'observed':gp,'comparison':rev,'variance':gp/rev*100 if rev else None,'unit':'GBP','materiality':'HIGH' if gps and abs(gp/gps)>=D('.1') else 'INFORMATIONAL','evidence':f'Annual customer summary evidences revenue {rev} and gross profit/Contribution 0 {gp}. No cost-to-serve allocation is inferred.','primitive':None,'lineage':[('WORKBOOK_ROW',f'{sh}:{row}','aggregate customer economics')]})
            results['CUS-02']=_execute(con,run,client,'CUS-02','AGGREGATE_CUSTOMER_MARGIN','PARTIAL-A','COMPLETED' if sig else 'NOT_RUN',sig,'Contribution 0 only; CTS unavailable')
        _replace_execution(con,run,'CUS-07'); _elig(con,run,client,'CUS-07','AGGREGATE_CUSTOMER_DISTRIBUTION','PARTIAL-A','Annual summary supports revenue Pareto/distribution only')
        cum=D('0'); n80=0
        for x in shares:
            if total and cum/total<D('.8'): cum+=x[1]; n80+=1
        sig=[{'type':'CUSTOMER_PARETO_DISTRIBUTION','observed':D(n80),'comparison':D(len(usable)),'variance':cum/total*100 if total else None,'unit':'COUNT','materiality':'INFORMATIONAL','evidence':f'{n80} of {len(usable)} customers account for at least 80% of captured annual revenue. This describes economic distribution, not customer quality.','primitive':None,'lineage':[('WORKBOOK_SHEET',shares[0][3],'aggregate customer summary')]}]
        results['CUS-07']=_execute(con,run,client,'CUS-07','AGGREGATE_CUSTOMER_DISTRIBUTION','PARTIAL-A','COMPLETED',sig,'Annual snapshot only')
    # Annual supplier summaries support spend concentration/dependency only.
    sus=[]
    for r in suppliers:
        if r.get('entity.supplier') and isinstance(r.get('metric.purchases'),(int,float)): sus.append((str(r['entity.supplier']),D(r['metric.purchases']),r['source_sheet'],r['source_row']))
    if sus:
        total=sum((x[1] for x in sus),D('0')); sus.sort(key=lambda x:x[1],reverse=True)
        for tid in ('SUP-01','SUP-03'): _replace_execution(con,run,tid); _elig(con,run,client,tid,'AGGREGATE_SUPPLIER_ANNUAL','PARTIAL-A','Annual supplier spend supports concentration; no item-level PPV, duplicate or disruption probability inferred')
        sig1=[{'type':'SUPPLIER_SPEND','entity_type':'SUPPLIER','entity_id':n,'observed':v,'comparison':v/total*100 if total else None,'unit':'GBP','materiality':'HIGH' if total and v/total>=D('.2') else 'INFORMATIONAL','evidence':f'Annual supplier summary evidences spend {v}, {v/total*100 if total else 0}% of captured spend. High spend alone is not overspend.','primitive':None,'lineage':[('WORKBOOK_ROW',f'{sh}:{row}','aggregate supplier summary')]} for n,v,sh,row in sus]
        results['SUP-01']=_execute(con,run,client,'SUP-01','AGGREGATE_SUPPLIER_ANNUAL','PARTIAL-A','COMPLETED',sig1,'Annual aggregate spend only')
        sig3=[{'type':'SUPPLIER_DEPENDENCY','entity_type':'SUPPLIER','entity_id':n,'observed':v/total*100,'comparison':v,'unit':'PERCENT','materiality':'HIGH' if v/total>=D('.3') else 'MEDIUM','evidence':f'{n} represents {v/total*100}% of captured annual supplier spend. This is concentration exposure only; no disruption probability or expected loss is inferred.','primitive':None,'lineage':[('WORKBOOK_ROW',f'{sh}:{row}','aggregate supplier summary')]} for n,v,sh,row in sus[:5] if total]
        results['SUP-03']=_execute(con,run,client,'SUP-03','AGGREGATE_SUPPLIER_ANNUAL','PARTIAL-A','COMPLETED',sig3,'Dependency based on spend concentration only; criticality/single-source status unavailable')
    # Work-centre summaries: accept plausible hour-based capacity; refuse internally impossible source units.
    cap=[]; refused=[]
    for r in ops:
        pc=r.get('metric.practical_capacity_hours'); used=r.get('metric.actual_productive_hours'); name=r.get('entity.work_centre')
        if name and isinstance(pc,(int,float)) and isinstance(used,(int,float)) and pc>0:
            ratio=D(used)/D(pc)*100
            (refused if ratio>D('250') else cap).append((str(name),D(pc),D(used),ratio,r['source_sheet'],r['source_row']))
    if cap or refused:
        _replace_execution(con,run,'PEO-04')
        if refused and not cap:
            _elig(con,run,client,'PEO-04','AGGREGATE_WORK_CENTRE_CAPACITY','UNAVAILABLE','Capacity source is internally inconsistent/implausible; utilisation refused rather than repaired')
            results['PEO-04']=_execute(con,run,client,'PEO-04','AGGREGATE_WORK_CENTRE_CAPACITY','UNAVAILABLE','NOT_RUN',[],'Capacity source is internally inconsistent/implausible; source units require confirmation')
        else:
            _elig(con,run,client,'PEO-04','AGGREGATE_WORK_CENTRE_CAPACITY','PARTIAL-A','Work-centre capacity supports operational utilisation; it is not employee productivity or a cash saving')
            sig=[{'type':'WORK_CENTRE_UTILISATION','entity_type':'WORK_CENTRE','entity_id':n,'observed':ratio,'comparison':pc,'variance':max(pc-used,D('0')),'unit':'PERCENT','materiality':'HIGH' if ratio>D('100') else 'INFORMATIONAL','evidence':f'{n}: practical capacity {pc} hours, actual productive hours {used}, utilisation {ratio}%. Over/unused capacity is operational evidence only; no financial saving or causal conclusion is inferred.','primitive':None,'lineage':[('WORKBOOK_ROW',f'{sh}:{row}','aggregate capacity summary')]} for n,pc,used,ratio,sh,row in cap]
            results['PEO-04']=_execute(con,run,client,'PEO-04','AGGREGATE_WORK_CENTRE_CAPACITY','PARTIAL-A','COMPLETED',sig,'Work-centre summary; no causal attribution')
    return {'customer_rows':len(usable),'supplier_rows':len(sus),'capacity_rows_accepted':len(cap),'capacity_rows_refused':len(refused),'tests':results}
